# -*- coding: utf-8 -*-
"""生成 issue #497–#504 批次交付状态（一页可打印 docx，含 rebase 与 PR）。"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DOC_PATH = r"C:\Users\wangj.KEVIN\projects\paleo_workstation\artifacts\issue-批次交付-2026-10-16.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.1); sec.bottom_margin = Cm(1.1)
sec.left_margin = Cm(1.3); sec.right_margin = Cm(1.3)

style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"; style.font.size = Pt(8)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
style.paragraph_format.space_after = Pt(1.5)
style.paragraph_format.space_before = Pt(0)


def set_cell(cell, text, bold=False, size=7.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text))
    run.font.size = Pt(size); run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_table(rows, widths, size=7.5):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            set_cell(cells[ci], val, bold=(ri == 0), size=size)
            if widths:
                cells[ci].width = Cm(widths[ci])
    return t


def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1.5)
    r = p.add_run(text); r.bold = True; r.font.size = Pt(8.5)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def para(text, size=8):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text); r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("issue #497–#504 批次交付状态")
r.bold = True; r.font.size = Pt(12.5)
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
para("PR #506 ｜ 分支 codex/ci-green-20261016 ｜ 基线 origin/master @ c720aef5（rebase 后）｜ 6 commits 已 push", size=7.5)

h1("一、提交（rebase 后，已 push）")
add_table([
    ["hash", "issue", "摘要"],
    ["08e74e24", "#504", ".gitignore 补 7 类规则 + 5 份 goal-loop ledger 归位"],
    ["64cf05a4", "附带", "predictionworkflow 直接 include arearules.h（ORT=OFF 构建必断 / C2653）"],
    ["c2091992", "#499", "CI Windows 两真红：异步清理竞态（QTRY）+ OGR 句柄（DeferredDelete）"],
    ["f37b3503", "#501", "环境债名单 reviewed_utc 强制 + 超期 warning + flaky 放行台账"],
    ["91ddd17a", "#502", "新闸 tools/check_ui_size.py（src/ui ≤1000 行，只缩不涨）+ baseline 12 条"],
    ["f82ba7d2", "#498+#502", "asan 磁盘水位闸（与 #505 叠加）+ 规模闸接线"],
], widths=[2.0, 1.8, 12.6])

h1("二、rebase 时的重叠处置（推送前必做的一步）")
add_table([
    ["我的提交", "master 已合入", "处置", "理由"],
    ["#500 child 改引用", "#495 同解（同一行）", "丢弃冗余提交", "问题已修复，重复改同一行只会制造冲突"],
    ["#498 清理 + -g1", "#505 清理更全 + -gz", "合并：采用 -gz，弃 -g1", "-gz 压缩调试段、信息不减；-g1 会丢变量/内联段"],
    ["—", "—", "保留我独有项", "Disk space guard（<40GB 早失败报明确原因）+ hostedtoolcache/node_modules"],
], widths=[2.8, 3.0, 3.4, 7.2])

h1("三、验证证据（新基线上复验，本机 Windows localdeps / Qt 6.11.2）")
add_table([
    ["项", "结果"],
    ["tst_mappingworkbench::removeDraftDeletesWorkingCopy", "3 passed / 0 failed ✅"],
    ["tst_versions::savedBytesAndPdfSurviveWorkingChanges", "3 passed / 0 failed ✅"],
    ["全量对照实验（stash 改动 vs 带改动）", "基线红 savedBytes → 修复后该用例转绿"],
    ["源码门禁 9 项（layering / ui_invariants / ui_size+selftest / i18n / env_redset validate+selftest / ci_scripts / docs_drift）", "9/9 绿"],
    ["check_ui_size --strict", "392 文件 / >1000 行 12 个 = baseline 12 条，clean"],
    ["YAML / CI 脚本引用", "可解析 + 15 个脚本引用全部存在"],
], widths=[8.6, 7.8])

h1("四、#503 的 18× 差距已定位（修正原判断）")
para("不是「Windows 慢 18 倍」，而是 **并行度 8× × 单测试 2.2×**：linux 走 ctest -j min(核数,8)（paleo-dev:104-112），"
     "而 paleo-dev.ps1:388 **未传 -j → 串行**。等效单进程：linux ≈147.6×8≈1181s vs Windows 2635s，剩余 2.2× 才是单测试本身差异。"
     "「Windows 必须串行」的理由（QSettings 注册表）在方向 81 落地测试私有注册表 hive 后可能已失效，ci.yml 注释疑为未同步更新。"
     "若并行可行，Windows 墙钟可从 44min 压到 ~11min —— 远优于拆巨型测试文件（拆文件不减少总墙钟）。已写入 #503，建议单独 PR 试跑 -j4（本机只有 2 个测试目标，无法本地验证并行安全）。")

h1("五、issue 收口")
add_table([
    ["issue", "状态"],
    ["#500", "已关闭（归属 #495，本批丢弃重复提交）"],
    ["#499", "已更新：2/3 处修复经 PR #506 提交；第三处（夹具清理缺陷）另立"],
    ["#497", "已更新进展表；待你拍板 master 健康度底线（是否恢复 strict required checks）"],
    ["#503", "已更新 18× 定位与拆分调研结论"],
    ["#498 / #501 / #502 / #504", "PR #506 中提交，待 CI 验证后自动关闭"],
], widths=[3.6, 12.8])

h1("六、下一步（你已选：拆巨型测试文件 + 定位 18× 差距）")
add_table([
    ["项", "状态", "说明"],
    ["定位 18× 差距", "✅ 已完成", "结论见第四节；实施（Windows 试跑 -j4）建议独立 PR，因需 CI 侧验证并行安全"],
    ["拆 tst_panels", "待实施", "5,337 行 = 单类 TestPanels 144 用例，命名已按页面分族（dataPage_* / predictPage_* / constraintPage_* / composePage_* / validatePage_* / folderConfirm_* / mainWindow_*）；需先抽共享 Fixture 到 tests/support/，属独立工程与独立分支/PR"],
], widths=[3.0, 2.2, 11.2])

doc.save(DOC_PATH)
print("saved:", DOC_PATH)
