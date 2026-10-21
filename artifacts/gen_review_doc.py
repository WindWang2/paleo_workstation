# -*- coding: utf-8 -*-
"""生成 paleo_workstation 整体性评价报告（一页可打印 docx）。"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DOC_PATH = r"C:\Users\wangj.KEVIN\projects\paleo_workstation\artifacts\paleo-整体评价-2026-10-16.docx"

doc = Document()

# 页面：A4 + 窄边距
sec = doc.sections[0]
sec.top_margin = Cm(1.2)
sec.bottom_margin = Cm(1.2)
sec.left_margin = Cm(1.5)
sec.right_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = "Microsoft YaHei"
style.font.size = Pt(8.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
style.paragraph_format.space_after = Pt(2)
style.paragraph_format.space_before = Pt(0)


def set_cell(cell, text, bold=False, size=8, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align


def add_table(rows, widths, header=True, size=8):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            set_cell(cells[ci], val, bold=(header and ri == 0), size=size)
            if widths:
                cells[ci].width = Cm(widths[ci])
    return t


def h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def para(text, size=8.5, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


# 标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(3)
r = p.add_run("paleo_workstation 整体性评价")
r.bold = True
r.font.size = Pt(13)
r.font.name = "Microsoft YaHei"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

para("基线 master @ ddd1325e（2026-10-16）｜评价范围：架构纪律 · 规模结构 · 测试与 CI 可信度 · 文档沉淀 · 工作区卫生",
     size=8, bold=False)

# 结论
h1("一、总体结论")
para("工程纪律是强项，验证链路是断点。分层契约（层标签 1103/1103、check_layering --strict 绿）与文档/账本沉淀（164 篇 docs + 92 份 ledger）"
     "达到可审计水准，src 内零 TODO/FIXME 说明技术债全部外置可追踪；但 master 已连续 10 次 push 的 CI 全红（windows / lint / linux-asan 三车道各有硬红）"
     "却仍持续合入，CI 事实上已从「门禁」退化为「日志」。次级风险是视图层规模失衡（113,101 行 / 38.6%、43 个 >1000 行文件且无防增长闸）"
     "与测试侧膨胀（184k 行测试码、Windows 全量 44 分钟）。建议本轮先修绿验证链路，再谈功能扩张。")

# 规模快照
h1("二、规模与结构快照")
add_table([
    ["指标", "数值", "指标", "数值"],
    ["src 代码", "293,344 行 / 1,103 文件", "注册 ctest", "406 项（win 388 / linux 373）"],
    ["tests 代码", "184,033 行 / 389 文件", "环境债豁免", "r0 26 + r1 22 = 48 条"],
    ["CMake 布线", "主 1,548 行 + 51 个 *.cmake", "层标签覆盖 / 严格分层", "1103/1103 ✅ / --strict 绿 ✅"],
    ["src 内 TODO/FIXME", "0（债务外置可追踪）✅", "docs / ledger / prompts", "164 / 92（5 未入库）/ 95"],
    ["巨型文件", ">1000 行 43 个、>1500 行 5 个", "build 体积 / 未跟踪文件", "85 GB / 23 个"],
], widths=[2.6, 5.4, 3.6, 4.4])

para("模块分布：ui 113,101（38.6%）· workflow 45,645 · algorithms 36,704 · io 24,830 · services 20,701 · qgis 18,328 · "
     "domain 11,900 · catalog 6,066 · ai 5,721 · metadata 5,419 · app 3,636 · linkage 876 · selfcheck 417。"
     "TOP5 巨型文件除 segyreader.cpp(1515) 外全在 ui/：seismic3dviewpanel 1811 · entitypanel 1620 · wellsectionpanel 1541 · mappingworkbenchpage 1517。",
     size=7.5)

# 六维评分
h1("三、六维健康度")
add_table([
    ["维度", "评级", "依据（实证）"],
    ["架构分层纪律", "✅ A", "层标签 1103/1103；check_layering --strict 绿；docs_drift 三项绿；UI_LAYER_PLAN 契约生效"],
    ["文档与账本沉淀", "✅ A-", "docs 164 篇 + 67 篇 progress + 92 份 goal-loop ledger，逐方向有交付证据；drift 检查仅 3 项（覆盖窄）"],
    ["测试覆盖规模", "⚠️ B", "406 测试 / 18.4 万行测试码；但 48 条豁免 + 非崩溃失败「串行重跑通过即放行」"],
    ["模块规模控制", "⚠️ C", "ui 占 38.6%；43 个 >1000 行；tst_panels 单文件 5,337 行；无防增长闸（参照 #474 同族缺口）"],
    ["工作区 / 仓库卫生", "⚠️ C", "build 85 GB；23 个未跟踪文件（*.log、调试 .ps1、osgeo4w-setup.exe）；5 份 ledger 未入库"],
    ["CI 可信度", "❌ D", "master 最近 10 次 push 全 failure；三车道各自硬红（详见第四节）"],
], widths=[2.6, 1.5, 11.9])

# 关键证据
h1("四、关键证据（run 38835970978 / 38756099683，均对应已合入 master 的提交）")
add_table([
    ["车道", "结论", "硬证据"],
    ["windows", "❌ 388 项 2 红", "tst_mappingworkbench::removeDraftDeletesWorkingCopy：'!QDir(workDir).exists()' 返回 FALSE；"
     "tst_versions::savedBytesAndPdfSurviveWorkingChanges：'QFile::remove(gpkg)' 返回 FALSE。env_redset 判「不在环境债名单 → 不放行」"],
    ["lint", "❌ clang-tidy", "src/workflow/mappingworkbench_impact.cpp:430:20 performance-unnecessary-copy-initialization（-warnings-as-errors）；"
     "跟进修复 PR 亦 failure"],
    ["linux-asan", "❌ 链接失败", "Free space left: 27 MB → ld.bfd: final link failed: No space left on device（2,231 链接目标，tst_wellcomposite / tst_smoke_realdata）"],
    ["linux", "✅（含金量存疑）", "373 项 2 红（tst_wellsiting_kernel、tst_evidencee2e）→ 串行重跑通过后 job 判绿，仅留 warning"],
    ["节奏", "⚠️", "10-06 峰值 191 commits/日 → 10-12~10-16 降至 3~16/日（由功能狂飙转入 review/收口）；26 个 open issue（#460–#485）待偿"],
], widths=[2.2, 2.6, 11.2])

# 建议
h1("五、处置建议（按优先级）")
para("1）先把 CI 修绿并恢复门禁效力：asan 车道降规模/换 runner、lint 修 clang-tidy 报错、Windows 定位句柄未释放；在三条车道恢复前，"
     "合并策略需显式登记红灯分诊 owner。2）给 48 条豁免名单加「到期复审、只减不增」机制，flaky 重跑通过须入台账。3）把「UI 单文件 ≤1000 行」"
     "做成防增长闸（与 #474 同族）。4）仓库卫生：补 .gitignore（*.log / 调试脚本 / 安装器），未入库 ledger 归位。", size=8)

# issue 候选
h1("六、拟提交 issue 清单（待确认）")
add_table([
    ["编号", "级", "标题", "一句话证据"],
    ["#497", "P0", "CI 门禁功能性失效：master 连续 10+ 次 push 全红仍持续合入", "run 38830506051…38835970978 十连 failure；AGENTS.md 既定流程以 --admin 绕过 required checks"],
    ["#498", "P1", "linux-asan 车道结构性磁盘耗尽，内存类检测形同虚设", "Free space 27 MB、2,231 链接目标、final link failed: No space left on device"],
    ["#499", "P1", "Windows 2 项真红疑为句柄未释放（工作副本/GeoPackage 删不掉）", "tst_mappingworkbench::removeDraftDeletesWorkingCopy + tst_versions::savedBytesAndPdfSurviveWorkingChanges"],
    ["#500", "P1", "lint 车道红于 mappingworkbench_impact.cpp:430 阻塞 master", "clang-tidy performance-unnecessary-copy-initialization（-warnings-as-errors）"],
    ["#501", "P2", "绿灯通胀：48 条豁免 + flaky 串行重跑放行，无到期复审", "r0 26 + r1 22；ci_ctest_gate.sh 对非崩溃失败重跑通过即放行"],
    ["#502", "P2", "视图层规模失衡：113k 行 / 38.6%、43 个 >1000 行，无防增长闸", "ui 113,101 行；TOP5 巨型文件 4 个在 ui/"],
    ["#503", "P2", "测试侧膨胀与反馈周期：tst_panels 5,337 行、全量 ctest 44 分钟", "Windows 全量 2,635 s；ui 845 s*proc、workflow 922 s*proc"],
    ["#504", "P3", "工作区卫生：23 未跟踪文件、.gitignore 无 *.log、5 份 ledger 未入库（补 #484）", "aqtinstall/launch/selfcheck.log、tools/*.ps1、vendor/osgeo4w-setup.exe"],
], widths=[1.1, 0.9, 6.8, 8.2])

doc.save(DOC_PATH)
print("saved:", DOC_PATH)
